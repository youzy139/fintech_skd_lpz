from __future__ import annotations

import csv
import json
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook

from market_core.domain.sources import SourceFile, SourceRecord

csv.field_size_limit(2**31 - 1)


def discover_sources(data_root: Path) -> list[SourceFile]:
    files: list[SourceFile] = []
    asset_root = data_root / "目标资产价格变量"
    files.extend(SourceFile.create(data_root, path, "asset") for path in sorted(asset_root.glob("*.csv")))
    excel_kinds = {"指标集1_数据.xlsx": "indicator", "商品期货日成交额.xlsx": "commodity", "日历数据.xlsx": "calendar", "公众号.xlsx": "publication"}
    files.extend(SourceFile.create(data_root, data_root / name, kind) for name, kind in excel_kinds.items())
    files.append(SourceFile.create(data_root, data_root / "科技.csv", "social"))
    files.append(SourceFile.create(data_root, data_root / "新闻" / "无标题.csv", "news"))
    research_root = data_root / "全球资金信息摘要" / "上科大版本摘要"
    files.extend(SourceFile.create(data_root, path, "research") for path in sorted(research_root.glob("*.md")))
    return sorted(files, key=lambda item: item.path)


def _date(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value)[:10]


def _row_payload(headers: list[str], row: tuple[Any, ...]) -> dict[str, Any]:
    return {header: value for header, value in zip(headers, row)}


def iter_excel(source: SourceFile) -> Iterator[SourceRecord]:
    workbook = load_workbook(source.absolute_path, read_only=True, data_only=True)
    index = 0
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            headers = [str(value) if value is not None else f"column_{position + 1}" for position, value in enumerate(header_row)]
            for row in rows:
                if not any(value is not None for value in row):
                    continue
                payload = _row_payload(headers, row)
                if source.kind == "indicator":
                    payload = {"_sheet": sheet.title, **payload}
                event_date = _date(payload.get("日期") or payload.get("OccurDate"))
                title = str(payload.get("IndicatorName") or payload.get("文章标题") or f"{sheet.title} · {index + 1}")
                yield SourceRecord(source.id, index, event_date, title, payload)
                index += 1
    finally:
        workbook.close()


def iter_csv(source: SourceFile) -> Iterator[SourceRecord]:
    with source.absolute_path.open(encoding="utf-8-sig", newline="") as handle:
        if source.kind == "news":
            reader = csv.reader(handle)
            for index, row in enumerate(reader):
                if len(row) != 11:
                    raise ValueError(f"新闻记录 {index + 1} 字段数为 {len(row)}，预期 11")
                payload = {f"column_{position + 1}": value for position, value in enumerate(row)}
                yield SourceRecord(source.id, index, _date(row[9]), row[1], payload)
            return
        reader = csv.DictReader(handle)
        for index, row in enumerate(reader):
            payload = dict(row)
            event_date = _date(payload.get("date") or payload.get("created_at_twitter"))
            title = str(payload.get("名称") or payload.get("text") or f"{source.absolute_path.stem} · {index + 1}")
            yield SourceRecord(source.id, index, event_date, title[:300], payload)


def iter_markdown(source: SourceFile) -> Iterator[SourceRecord]:
    parts = source.absolute_path.stem.split("-", 3)
    event_date = "-".join(parts[:3]) if len(parts) == 4 else None
    institution = parts[3].split("-", 1)[0] if len(parts) == 4 else ""
    title = parts[3].split("-", 1)[1] if len(parts) == 4 and "-" in parts[3] else source.absolute_path.stem
    text = source.absolute_path.read_text(encoding="utf-8-sig")
    payload = {"institution": institution, "title": title, "content": text}
    yield SourceRecord(source.id, 0, event_date, title, payload)


def iter_records(source: SourceFile) -> Iterator[SourceRecord]:
    if source.kind in {"indicator", "commodity", "calendar", "publication"}:
        return iter_excel(source)
    if source.kind in {"asset", "news", "social"}:
        return iter_csv(source)
    if source.kind == "research":
        return iter_markdown(source)
    raise ValueError(f"未知数据类型: {source.kind}")


def source_sheets(source: SourceFile) -> list[str]:
    if source.kind not in {"indicator", "commodity", "calendar", "publication"}:
        return []
    workbook = load_workbook(source.absolute_path, read_only=True, data_only=True)
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


def json_payload(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, default=str)
